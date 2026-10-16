from __future__ import annotations

import zipfile
from datetime import datetime

import pytest
from openpyxl import load_workbook

from food_agent import config
from food_agent.ledger import HEADERS, Ledger, LedgerBusyError, LedgerError, LedgerRow, neutralize_formula

D1 = datetime(2026, 10, 15, 12, 30, 45)
D2 = datetime(2026, 10, 16, 8, 0, 0)


def row(when=D1, uploader="小明", name="牛肉麵", kcal=650, weight=550, photos=("photos/a.jpg",)):
    return LedgerRow(when, uploader, name, "麵條、牛肉", weight, kcal, list(photos))


def sheet(path):
    return load_workbook(path).worksheets[0]


def test_creates_file_with_headers(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    assert ledger.append([row()]) == 650
    ws = sheet(tmp_path / "log.xlsx")
    assert [c.value for c in ws[1]] == HEADERS
    assert [c.value for c in ws[2]][:7] == [D1, "小明", "牛肉麵", "麵條、牛肉", 550, 650, 650]
    assert ws.cell(2, 8).value == "photos/a.jpg"


def test_appends_to_the_end_without_overwriting(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(name="第一餐")])
    ledger.append([row(name="第二餐")])
    ledger.append([row(name="第三餐")])
    ws = sheet(tmp_path / "log.xlsx")
    assert ws.max_row == 4
    assert [ws.cell(r, 3).value for r in (2, 3, 4)] == ["第一餐", "第二餐", "第三餐"]


def test_daily_total_is_whole_channel_not_per_person(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    assert ledger.append([row(uploader="小明", kcal=650)]) == 650
    assert ledger.append([row(uploader="小華", kcal=400)]) == 1050  # 不分人
    assert ledger.append([row(uploader="小明", kcal=100)]) == 1150
    ws = sheet(tmp_path / "log.xlsx")
    assert [ws.cell(r, 7).value for r in (2, 3, 4)] == [650, 1050, 1150]


def test_daily_total_resets_on_new_day(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(when=D1, kcal=650)])
    assert ledger.append([row(when=D2, kcal=300)]) == 300


def test_midnight_boundary(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(when=datetime(2026, 10, 15, 23, 59, 59), kcal=500)])
    assert ledger.append([row(when=datetime(2026, 10, 16, 0, 0, 0), kcal=100)]) == 100


def test_total_uses_hand_edited_calories(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(kcal=650)])
    wb = load_workbook(tmp_path / "log.xlsx")
    wb.worksheets[0].cell(2, 6).value = 100  # 你手動把熱量改成 100
    wb.save(tmp_path / "log.xlsx")
    assert ledger.append([row(kcal=50)]) == 150


def test_non_food_rows_do_not_add_calories(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(kcal=650)])
    non_food = LedgerRow(D1, "小明", "不是食物", "一隻貓", None, None, ["photos/cat.jpg"])
    assert ledger.append([non_food]) == 650
    ws = sheet(tmp_path / "log.xlsx")
    assert [c.value for c in ws[3]][2:7] == ["不是食物", "一隻貓", None, None, 650]


def test_multiple_rows_from_one_message_share_the_total(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    non_food = LedgerRow(D1, "小明", "不是食物", "一隻貓", None, None, ["photos/cat.jpg"])
    assert ledger.append([row(kcal=900), non_food]) == 900
    ws = sheet(tmp_path / "log.xlsx")
    assert ws.cell(2, 7).value == 900 and ws.cell(3, 7).value == 900


def test_multiple_photos_in_one_cell(tmp_path):
    ledger = Ledger(tmp_path / "log.xlsx")
    ledger.append([row(photos=("photos/a.jpg", "photos/b.jpg"))])
    assert sheet(tmp_path / "log.xlsx").cell(2, 8).value == "photos/a.jpg\nphotos/b.jpg"


# ---- 防護 4 ----
@pytest.mark.parametrize("evil", ["=1+1", "+cmd|' /C calc'!A0", "-2", "@SUM(1)", '=HYPERLINK("http://x","點我")'])
def test_formula_text_is_stored_as_text(tmp_path, evil):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row(uploader=evil, name=evil)])
    ws = sheet(path)
    cell = ws.cell(2, 3)
    assert cell.data_type == "s"
    assert cell.value == "'" + evil
    # xlsx 內部的檔案裡完全沒有公式標籤
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in xml and "<f " not in xml


def test_neutralize_formula_leaves_normal_text_alone():
    assert neutralize_formula("牛肉麵") == "牛肉麵"
    assert neutralize_formula("") == ""
    assert neutralize_formula("a=b") == "a=b"
    assert neutralize_formula("\t=1") == "'\t=1"


# ---- 容錯 ----
def test_header_mismatch_refuses_to_write(tmp_path):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row()])
    wb = load_workbook(path)
    wb.worksheets[0].cell(1, 1).value = "被改過的標題"
    wb.save(path)
    with pytest.raises(LedgerError):
        Ledger(path).append([row()])
    assert sheet(path).max_row == 2  # 沒有被寫入


def test_corrupted_file_raises_ledger_error_and_is_kept(tmp_path):
    path = tmp_path / "log.xlsx"
    path.write_bytes(b"this is not a zip")
    with pytest.raises(LedgerError):
        Ledger(path).append([row()])
    assert path.read_bytes() == b"this is not a zip"


def test_locked_file_warns_retries_and_does_not_crash(tmp_path, monkeypatch, caplog):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row()])
    monkeypatch.setattr(config, "EXCEL_RETRY_TIMES", 3)
    monkeypatch.setattr(config, "EXCEL_RETRY_WAIT_S", 0)
    calls = []

    def locked(src, dst):
        calls.append(1)
        raise PermissionError("[WinError 32] 另一個程序正在使用此檔案")

    monkeypatch.setattr("food_agent.ledger.os.replace", locked)
    with pytest.raises(LedgerBusyError):
        Ledger(path).append([row()])
    assert len(calls) == 3
    warnings = [r.getMessage() for r in caplog.records if r.levelname == "WARNING" and "[警告]" in r.getMessage()]
    assert len(warnings) == 3 and "1/3" in warnings[0] and "3/3" in warnings[2]
    assert sheet(path).max_row == 2  # 原檔沒被動到
    assert not list(tmp_path.glob("*.tmp.xlsx"))  # 暫存檔有清掉


def test_succeeds_after_file_is_released(tmp_path, monkeypatch):
    path = tmp_path / "log.xlsx"
    Ledger(path).append([row()])
    monkeypatch.setattr(config, "EXCEL_RETRY_WAIT_S", 0)
    import os

    real_replace = os.replace
    state = {"n": 0}

    def flaky(src, dst):
        state["n"] += 1
        if state["n"] < 3:
            raise PermissionError("locked")
        return real_replace(src, dst)

    monkeypatch.setattr("food_agent.ledger.os.replace", flaky)
    assert Ledger(path).append([row(kcal=100)]) == 750
    assert sheet(path).max_row == 3
