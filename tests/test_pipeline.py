from __future__ import annotations

import asyncio

from conftest import food_analysis, make_job, non_food_analysis
from food_agent import config
from food_agent.ledger import Ledger
from food_agent.pipeline import process_job
from food_agent.schema import Analysis, AnalysisError, Meal, PhotoVerdict
from openpyxl import load_workbook


def _run(job, settings, analysis=None, error=None):
    async def fake_analyze(photos, caption, s):
        if error:
            raise error
        return analysis

    ledger = Ledger(settings.ledger_path)
    return asyncio.run(process_job(job, settings, ledger, analyze_fn=fake_analyze))


def _rows(settings):
    return list(load_workbook(settings.ledger_path).worksheets[0].iter_rows(min_row=2, values_only=True))


def _photos(settings):
    return sorted(p.name for p in settings.photos_dir.iterdir())


def test_food_flow(settings):
    out = _run(make_job(), settings, food_analysis())
    assert out.ok
    assert _photos(settings) == ["2026-10-15 12-30-45 牛肉麵.jpg"]
    (row,) = _rows(settings)
    assert row[1:8] == ("小明", "牛肉麵", "麵條、牛肉", 550, 650, 650, "photos/2026-10-15 12-30-45 牛肉麵.jpg")
    # 回覆有是不是食物、成份、重量、熱量，最後一定有當日已攝取熱量
    for needle in ("<@1>", "這是食物：牛肉麵", "成份：麵條、牛肉", "550", "650", "當日已攝取熱量：650 大卡"):
        assert needle in out.text
    assert out.text.rstrip().splitlines()[-1] != ""
    assert "當日已攝取熱量" in out.text


def test_daily_total_accumulates_across_people(settings):
    _run(make_job(uid=1, name="小明"), settings, food_analysis(kcal=650))
    out = _run(make_job(uid=2, name="小華"), settings, food_analysis(name="沙拉", kcal=200))
    assert "當日已攝取熱量：850 大卡" in out.text
    assert [r[6] for r in _rows(settings)] == [650, 850]


def test_non_food_is_recorded_saved_and_replied(settings):
    out = _run(make_job(), settings, non_food_analysis("馬克杯"))
    assert out.ok
    assert _photos(settings) == ["2026-10-15 12-30-45 不是食物_馬克杯.jpg"]
    (row,) = _rows(settings)
    assert row[2] == "不是食物" and row[3] == "馬克杯" and row[4] is None and row[5] is None and row[6] == 0
    assert "你傳的不是食物，是馬克杯。" in out.text
    assert "當日已攝取熱量：0 大卡" in out.text


def test_mixed_message_makes_two_rows(settings):
    analysis = Analysis(
        photos=(PhotoVerdict(1, True, ""), PhotoVerdict(2, True, ""), PhotoVerdict(3, False, "一隻貓")),
        meal=Meal("牛排套餐", ("牛排", "薯條"), 480, 920),
    )
    out = _run(make_job(n_photos=3), settings, analysis)
    assert out.ok
    rows = _rows(settings)
    assert [r[2] for r in rows] == ["牛排套餐", "不是食物"]
    assert rows[0][7].splitlines() == [
        "photos/2026-10-15 12-30-45 牛排套餐.jpg",
        "photos/2026-10-15 12-30-45 牛排套餐-2.jpg",
    ]
    assert [r[6] for r in rows] == [920, 920]
    assert "食物（第 1、2 張）：牛排套餐" in out.text and "第 3 張不是食物，是一隻貓。" in out.text
    assert len(_photos(settings)) == 3


def test_analysis_failure_leaves_no_trace(settings):
    out = _run(make_job(), settings, error=AnalysisError("boom"))
    assert not out.ok and out.text == "分析失敗，請再傳一次"
    assert _photos(settings) == []
    assert not settings.ledger_path.exists()


def test_unexpected_error_does_not_crash(settings):
    out = _run(make_job(), settings, error=RuntimeError("something odd"))
    assert not out.ok and out.text == "分析失敗，請再傳一次"
    assert _photos(settings) == []


def test_busy_excel_deletes_saved_photos_and_replies(settings, monkeypatch):
    monkeypatch.setattr(config, "EXCEL_RETRY_TIMES", 2)
    monkeypatch.setattr(config, "EXCEL_RETRY_WAIT_S", 0)

    def locked(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr("food_agent.ledger.os.replace", locked)
    out = _run(make_job(), settings, food_analysis())
    assert not out.ok and "記錄檔現在寫不進去" in out.text
    assert _photos(settings) == []  # 已存的照片有刪掉


def test_truncated_note_in_reply(settings):
    out = _run(make_job(truncated=True), settings, food_analysis())
    assert f"一次最多處理 {config.MAX_PHOTOS_PER_MESSAGE} 張" in out.text


def test_evil_names_from_claude(settings):
    # 防護 3 + 防護 4：claude 回來的名稱是惡意的
    evil = "=1+1/..\\..\\x:y"  # 開頭是 =（防護 4），中間有路徑亂跑的字（防護 3）
    out = _run(make_job(), settings, food_analysis(name=evil))
    assert out.ok
    files = [p for p in settings.data_dir.rglob("*") if p.is_file() and p != settings.ledger_path]
    assert files and all(p.parent == settings.photos_dir for p in files)
    ws = load_workbook(settings.ledger_path).worksheets[0]
    assert ws.cell(2, 3).data_type == "s" and ws.cell(2, 3).value.startswith("'")


def test_evil_nickname_goes_to_excel_as_text(settings):
    out = _run(make_job(name="=HYPERLINK(\"http://x\",\"點我\")"), settings, food_analysis())
    assert out.ok
    cell = load_workbook(settings.ledger_path).worksheets[0].cell(2, 2)
    assert cell.data_type == "s" and cell.value.startswith("'=")


def test_message_time_not_processing_time_decides_the_day(settings):
    from datetime import datetime

    _run(make_job(when=datetime(2026, 10, 15, 23, 59, 0)), settings, food_analysis(kcal=500))
    out = _run(make_job(when=datetime(2026, 10, 16, 0, 1, 0)), settings, food_analysis(kcal=100))
    assert "當日已攝取熱量：100 大卡" in out.text
