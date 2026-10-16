"""Excel 紀錄檔（docs/architecture.md 第 8 節）。

這個檔案裡有一段防護：
  防護 4：寫進 Excel 的文字若以 = + - @ 開頭，先處理成純文字
前面先放一行「沒有防護時的預設寫法」，防護段落再覆蓋它。
"""

from __future__ import annotations

import logging
import os
import time
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from food_agent import config

log = logging.getLogger("food_agent.ledger")

HEADERS = [
    "日期時間",
    "上傳者",
    "食物敘述",
    "成份",
    "重量(公克)",
    "熱量(大卡)",
    "當日已攝取熱量(大卡)",
    "照片存放位置",
]
_COL_KCAL = 6  # 熱量(大卡) 在第幾欄（從 1 開始）
_DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r", "\n")


class LedgerError(Exception):
    """紀錄檔有問題（標題對不上、檔案壞掉），重試也沒用。"""


class LedgerBusyError(LedgerError):
    """紀錄檔一直寫不進去（通常是你正用 Excel 開著它）。"""


@dataclass
class LedgerRow:
    when: datetime  # 本機時間，不帶時區
    uploader: str
    name: str  # 食物敘述；非食物寫「不是食物」
    ingredients: str  # 成份；非食物寫它是什麼
    weight_g: int | None
    kcal: int | None
    photo_paths: list[str] = field(default_factory=list)  # 相對於 DATA_DIR


def neutralize_formula(value: str) -> str:
    """開頭是 = + - @ 或控制字元時，在最前面加單引號。"""
    if value[:1] in _FORMULA_STARTS:
        return "'" + value
    return value


def _set_text(cell, value: str) -> None:
    # ---- 沒有防護 4 時的預設：原文寫入（openpyxl 會把 = 開頭的字當成公式）----
    cell.value = value
    # ==== 防護 4：Excel 公式字元處理　開始 ====
    cell.value = neutralize_formula(value)  # 以 = + - @ 開頭的，前面加單引號
    cell.data_type = "s"  # 明確標成文字格，不讓它被當成公式
    # ==== 防護 4 結束 ====


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%Y-%m-%d %H:%M:%S").date()
        except ValueError:
            return None
    return None


def _as_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return None
    return None


class Ledger:
    def __init__(self, path: Path):
        self.path = path

    def append(self, rows: list[LedgerRow]) -> int:
        """把一則訊息產生的列加到最後，回傳寫入後「這一天整個頻道」的總熱量。
        會 sleep，要在執行緒裡呼叫（asyncio.to_thread）。"""
        if not rows:
            raise LedgerError("沒有要寫的列")
        for attempt in range(1, config.EXCEL_RETRY_TIMES + 1):
            try:
                return self._write_once(rows)
            except PermissionError as exc:
                msg = (
                    f"[警告] 紀錄檔 {self.path.name} 寫不進去（可能正在用 Excel 開著，請先關掉）。"
                    f"第 {attempt}/{config.EXCEL_RETRY_TIMES} 次：{exc}"
                )
                log.warning(msg)  # 終端機和 agent.log 都會有（見 logs.py）
                if attempt < config.EXCEL_RETRY_TIMES:
                    time.sleep(config.EXCEL_RETRY_WAIT_S)
        raise LedgerBusyError(f"紀錄檔 {self.path} 重試 {config.EXCEL_RETRY_TIMES} 次仍然寫不進去")

    def _write_once(self, rows: list[LedgerRow]) -> int:
        tmp = self.path.with_name(self.path.stem + ".tmp.xlsx")
        try:
            wb, ws = self._open()
            day = rows[0].when.date()
            before = self._day_total(ws, day)
            total = before + sum(r.kcal or 0 for r in rows)

            for row in rows:
                r = ws.max_row + 1
                c = ws.cell(row=r, column=1, value=row.when)
                c.number_format = _DATETIME_FORMAT
                _set_text(ws.cell(row=r, column=2), row.uploader)
                _set_text(ws.cell(row=r, column=3), row.name)
                _set_text(ws.cell(row=r, column=4), row.ingredients)
                ws.cell(row=r, column=5, value=row.weight_g)
                ws.cell(row=r, column=6, value=row.kcal)
                ws.cell(row=r, column=7, value=total)
                photos = ws.cell(row=r, column=8)
                _set_text(photos, "\n".join(row.photo_paths))
                photos.alignment = Alignment(wrap_text=True, vertical="top")

            wb.save(tmp)
            os.replace(tmp, self.path)  # 先存暫存檔再取代：寫到一半當掉，原檔不會變成半份
            return total
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass

    def _open(self):
        if self.path.exists():
            try:
                wb = load_workbook(self.path)
            except (zipfile.BadZipFile, KeyError) as exc:
                raise LedgerError(f"紀錄檔 {self.path} 打不開，可能已損壞：{exc}") from exc
            ws = wb.worksheets[0]
            header = [ws.cell(row=1, column=i + 1).value for i in range(len(HEADERS))]
            if header != HEADERS:
                raise LedgerError(
                    f"紀錄檔的標題列和預期不同，為了不弄壞你的檔案，這次不寫入。目前：{header}"
                )
            return wb, ws
        wb = Workbook()
        ws = wb.active
        ws.title = "紀錄"
        ws.append(HEADERS)
        for col, width in zip("ABCDEFGH", (20, 14, 18, 30, 12, 12, 22, 48)):
            ws.column_dimensions[col].width = width
        return wb, ws

    @staticmethod
    def _day_total(ws, day: date) -> int:
        """這一天所有列的「熱量」加總（用每一餐的熱量重新加，不是拿上一列的累計）。"""
        total = 0.0
        for values in ws.iter_rows(min_row=2, max_col=_COL_KCAL, values_only=True):
            if _as_date(values[0]) != day:
                continue
            kcal = _as_number(values[_COL_KCAL - 1])
            if kcal is not None:
                total += kcal
        return int(round(total))
